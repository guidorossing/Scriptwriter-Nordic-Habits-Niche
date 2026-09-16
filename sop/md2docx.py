import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0B, 0x4F, 0x75)
INK  = RGBColor(0x1A, 0x1A, 0x1A)

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10.5); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

for name, size, bold, color, before in (
    ('Heading 1', 19, True, NAVY, 20), ('Heading 2', 14, True, NAVY, 16),
    ('Heading 3', 11.5, True, RGBColor(0x2C,0x2C,0x2C), 12)):
    s = doc.styles[name]
    s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=color
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(5)

def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'),'clear'); el.set(qn('w:fill'),hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


from docx.opc.constants import RELATIONSHIP_TYPE as RT
LINK = RGBColor(0x0B, 0x4F, 0x75)

def add_hyperlink(par, text, url):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    for tag, val in (('w:color', '0B4F75'),):
        e = OxmlElement(tag); e.set(qn('w:val'), val); rPr.append(e)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); par._p.append(h)

MDLINK = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')

TOKEN = re.compile(r'(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`]+`)')
def rich(par, text):
    text = re.sub(r'<(https?://[^>]+)>', r'\1', text)
    pos = 0
    for m in MDLINK.finditer(text):
        _plain(par, text[pos:m.start()])
        add_hyperlink(par, m.group(1), m.group(2))
        pos = m.end()
    _plain(par, text[pos:])

def _plain(par, text):
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(9.5)
        elif part.startswith('*') and part.endswith('*'):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)

lines = open('SOP3-celebrity-thumbnails.md', encoding='utf-8').read().split('\n')
i = 0
while i < len(lines):
    ln = lines[i].rstrip()

    # --- tabel ---
    if ln.startswith('|') and i+1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            if not re.match(r'^\|[\s:|-]+\|$', lines[i].strip()):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Inches(6.5/ncol)]*ncol
        t.columns and None
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncol):
                cells[ci].width = widths[ci]
                p = cells[ci].paragraphs[0]; p.paragraph_format.space_after = Pt(2)
                rich(p, row[ci] if ci < len(row) else '')
                for r in p.runs: r.font.size = Pt(9.5)
                if ri == 0:
                    shade(cells[ci], 'E4ECF1')
                    for r in p.runs: r.bold = True
        doc.add_paragraph()
        continue

    if ln.startswith('# '):    rich(doc.add_paragraph(style='Heading 1'), ln[2:])
    elif ln.startswith('## '): rich(doc.add_paragraph(style='Heading 2'), ln[3:])
    elif ln.startswith('### '):rich(doc.add_paragraph(style='Heading 3'), ln[4:])
    elif ln.strip() == '---':
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        pb = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6'); b.set(qn('w:color'),'C8D0D6')
        pb.append(b); p._p.get_or_add_pPr().append(pb)
    elif ln.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
        rich(p, ln[2:])
    elif re.match(r'^\d+\. ', ln):
        p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(3)
        rich(p, re.sub(r'^\d+\. ', '', ln))
    elif ln.startswith('> '):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
        rich(p, ln[2:])
        for r in p.runs: r.italic = True
    elif ln.strip():
        rich(doc.add_paragraph(), ln)
    i += 1

out = 'SOP 3 - Thumbnails (Celebrity Clips).docx'
doc.save(out)
print("geschreven:", out)
